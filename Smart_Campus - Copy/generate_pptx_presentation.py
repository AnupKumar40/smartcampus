"""
generate_pptx_presentation.py
Generates a 16-slide high-end, professionally styled PowerPoint presentation
for Smart Campus: AI-Driven Campus ERP, Placement Readiness & ATS Resume Analytics System.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette Definitions
    BG_DARK = RGBColor(11, 19, 43)        # Deep Navy / Space Cadet (#0B132B)
    CARD_DARK = RGBColor(28, 37, 65)      # Slate Navy Card (#1C2541)
    CARD_BORDER = RGBColor(58, 80, 107)   # Muted Blue Border (#3A506B)
    TEXT_WHITE = RGBColor(255, 255, 255)
    TEXT_MUTED = RGBColor(203, 213, 225)  # Slate 300
    TEXT_ACCENT = RGBColor(56, 189, 248)  # Cyan 400 (#38BDF8)
    COLOR_ACCENT = RGBColor(56, 189, 248) # Cyan 400 (#38BDF8)
    COLOR_PRIMARY = RGBColor(37, 99, 235) # Blue 600
    COLOR_SUCCESS = RGBColor(16, 185, 129)# Emerald 500 (#10B981)
    COLOR_AMBER = RGBColor(245, 158, 11)  # Amber 500 (#F59E0B)
    COLOR_PURPLE = RGBColor(168, 85, 247) # Purple 500 (#A855F7)

    def set_slide_background(slide, color=BG_DARK):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category="SMART CAMPUS PROJECT PRESENTATION"):
        # Category / Breadcrumb
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.name = "Segoe UI"
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = TEXT_ACCENT

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.65))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Segoe UI"
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

    def add_card(slide, left, top, width, height, bg_color=CARD_DARK, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.fill.background()
        return card

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, RGBColor(7, 12, 28))

    # Decorative accent banner
    accent_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.1), Inches(1.5), Inches(0.08))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = TEXT_ACCENT
    accent_bar.line.fill.background()

    # Title & Subtitle box
    t_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(3.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "SMART CAMPUS"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(44)
    p0.font.bold = True
    p0.font.color.rgb = TEXT_WHITE

    p1 = tf1.add_paragraph()
    p1.text = "AI-Driven Campus ERP, Placement Readiness & ATS Resume Analytics Platform"
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(20)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_ACCENT
    p1.space_before = Pt(10)

    p2 = tf1.add_paragraph()
    p2.text = "A Comprehensive Machine Learning & Full-Stack System for Higher Education Institutions"
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(14)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_before = Pt(8)

    # 4 Tech badges at bottom
    badges = [
        ("ML Placement Engine", "Random Forest (76.3% Acc)"),
        ("NLP ATS Resume Parser", "TF-IDF & Cosine Similarity"),
        ("Academic ERP Hub", "Live Timetable & Attendance"),
        ("Corporate Placement Portal", "Job Drives & Application ATS")
    ]
    b_width = Inches(2.7)
    b_gap = Inches(0.28)
    for i, (b_title, b_sub) in enumerate(badges):
        bx = Inches(0.8) + i * (b_width + b_gap)
        add_card(slide1, bx, Inches(5.1), b_width, Inches(1.6), CARD_DARK, CARD_BORDER)
        bbox = slide1.shapes.add_textbox(bx + Inches(0.2), Inches(5.25), b_width - Inches(0.4), Inches(1.3))
        btf = bbox.text_frame
        btf.word_wrap = True
        bp0 = btf.paragraphs[0]
        bp0.text = b_title
        bp0.font.name = "Segoe UI"
        bp0.font.size = Pt(12)
        bp0.font.bold = True
        bp0.font.color.rgb = TEXT_WHITE

        bp1 = btf.add_paragraph()
        bp1.text = b_sub
        bp1.font.name = "Segoe UI"
        bp1.font.size = Pt(10)
        bp1.font.color.rgb = TEXT_ACCENT
        bp1.space_before = Pt(4)

    # =========================================================================
    # SLIDE 2: EXECUTIVE SUMMARY & PROJECT VISION
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "Executive Summary & Core Project Vision", "01. INTRODUCTION")

    # 3 High level cards
    cols = [
        ("1. Integrated Academic ERP", 
         "Consolidates daily college administration into a unified portal. Provides a real-time weekly schedule, class status indicators, subject-wise attendance thresholds, course catalogs, and faculty profiles.",
         "Academic Operational Excellence", COLOR_PRIMARY),
        ("2. Machine Learning Placement Engine", 
         "Leverages a Random Forest Classifier trained on 15,000 engineering student records. Evaluates 8 key predictive dimensions (DSA, CGPA, Communication, etc.) and provides instant tailored recommendations.",
         "Predictive Career Intelligence", COLOR_SUCCESS),
        ("3. Intelligent ATS Resume Analyzer", 
         "Parses PDF/DOCX resumes using PyPDF2 & python-docx. Extracts 80+ technical and soft skills, measures cosine similarity against job descriptions, and gives actionable feedback to maximize shortlisting.",
         "NLP-Driven Employability Boost", COLOR_PURPLE)
    ]
    c_w = Inches(3.68)
    c_g = Inches(0.33)
    for i, (head, desc, tag, col) in enumerate(cols):
        cx = Inches(0.8) + i * (c_w + c_g)
        add_card(slide2, cx, Inches(1.6), c_w, Inches(5.2), CARD_DARK, col)
        
        # Tag
        tbox = slide2.shapes.add_textbox(cx + Inches(0.3), Inches(1.85), c_w - Inches(0.6), Inches(0.4))
        ttf = tbox.text_frame
        tp = ttf.paragraphs[0]
        tp.text = tag.upper()
        tp.font.name = "Segoe UI"
        tp.font.size = Pt(10)
        tp.font.bold = True
        tp.font.color.rgb = col

        # Content
        cbox = slide2.shapes.add_textbox(cx + Inches(0.3), Inches(2.35), c_w - Inches(0.6), Inches(4.2))
        ctf = cbox.text_frame
        ctf.word_wrap = True
        cp0 = ctf.paragraphs[0]
        cp0.text = head
        cp0.font.name = "Segoe UI"
        cp0.font.size = Pt(15)
        cp0.font.bold = True
        cp0.font.color.rgb = TEXT_WHITE
        cp0.space_after = Pt(12)

        cp1 = ctf.add_paragraph()
        cp1.text = desc
        cp1.font.name = "Segoe UI"
        cp1.font.size = Pt(11)
        cp1.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 3: PROBLEM STATEMENT & MARKET GAPS
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "Problem Statement & Industry Gaps", "02. RESEARCH CONTEXT")

    problems = [
        ("Fragmented Campus ERPs", 
         "Existing collegiate software tools isolate attendance and marks into static silos. They offer zero correlation between daily student performance and eventual campus recruitment success."),
        ("Late Detection of Career Deficiencies", 
         "Students typically realize their lack of coding proficiency (DSA), low aptitude scores, or communication deficits during 7th/8th semester drives when remedial action is nearly impossible."),
        ("High Rejection Rates at Initial ATS Screening", 
         "Over 75% of engineering resumes are screened out before human review due to missing keywords, non-standard formatting, and misalignment with industry job descriptions."),
        ("Administrative Placement Overhead", 
         "Placement cells spend hundreds of hours manually verifying CGPA eligibility criteria, emailing drive updates, tracking interview pipelines, and managing recruiter logistics.")
    ]
    card_h = Inches(2.4)
    for i, (p_title, p_desc) in enumerate(problems):
        row = i // 2
        col = i % 2
        px = Inches(0.8) + col * (Inches(5.6) + Inches(0.5))
        py = Inches(1.6) + row * (card_h + Inches(0.35))
        add_card(slide3, px, py, Inches(5.6), card_h, CARD_DARK, CARD_BORDER)

        pbox = slide3.shapes.add_textbox(px + Inches(0.3), py + Inches(0.25), Inches(5.0), card_h - Inches(0.5))
        ptf = pbox.text_frame
        ptf.word_wrap = True
        
        p_h = ptf.paragraphs[0]
        p_h.text = f"❌  {p_title}"
        p_h.font.name = "Segoe UI"
        p_h.font.size = Pt(14)
        p_h.font.bold = True
        p_h.font.color.rgb = RGBColor(239, 68, 68) # Red 500
        p_h.space_after = Pt(8)

        p_b = ptf.add_paragraph()
        p_b.text = p_desc
        p_b.font.name = "Segoe UI"
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 4: PROJECT OBJECTIVES & KEY GOALS
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "Project Objectives & Technical Scope", "03. PROJECT SCOPE")

    objectives = [
        ("Unified Academic Management", "Provide seamless student attendance tracking, dynamic weekly scheduling with active-class detection, and academic profile transparency.", "GOAL 1"),
        ("Predictive Placement Analytics", "Implement machine learning classification to forecast student placement probability with actionable guidance across 8 critical skill dimensions.", "GOAL 2"),
        ("Intelligent ATS Resume Optimization", "Empower students to evaluate resumes against corporate job descriptions using TF-IDF vectorization, cosine similarity, and skill extraction.", "GOAL 3"),
        ("End-to-End Placement Recruitment ATS", "Facilitate corporate job listings, automated CGPA eligibility screening, student job applications, and transparent status tracking.", "GOAL 4"),
        ("Data-Driven Placement Cell Dashboard", "Provide institutional administrators with high-level KPIs, branch-wise statistics, and tools to manage drives and candidate statuses.", "GOAL 5")
    ]
    for i, (title, desc, badge) in enumerate(objectives):
        oy = Inches(1.55) + i * Inches(1.05)
        add_card(slide4, Inches(0.8), oy, Inches(11.73), Inches(0.92), CARD_DARK, CARD_BORDER)
        
        # Badge
        bbox = slide4.shapes.add_textbox(Inches(1.0), oy + Inches(0.2), Inches(1.2), Inches(0.5))
        btf = bbox.text_frame
        bp = btf.paragraphs[0]
        bp.text = badge
        bp.font.name = "Segoe UI"
        bp.font.size = Pt(11)
        bp.font.bold = True
        bp.font.color.rgb = TEXT_ACCENT

        # Text
        tbox = slide4.shapes.add_textbox(Inches(2.3), oy + Inches(0.12), Inches(10.0), Inches(0.7))
        ttf = tbox.text_frame
        ttf.word_wrap = True
        tp0 = ttf.paragraphs[0]
        tp0.text = title
        tp0.font.name = "Segoe UI"
        tp0.font.size = Pt(13)
        tp0.font.bold = True
        tp0.font.color.rgb = TEXT_WHITE
        
        tp1 = ttf.add_paragraph()
        tp1.text = desc
        tp1.font.name = "Segoe UI"
        tp1.font.size = Pt(10)
        tp1.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 5: SYSTEM ARCHITECTURE & HIGH LEVEL DESIGN
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "System Architecture & Layered Framework", "04. ARCHITECTURE")

    layers = [
        ("Layer 1: Presentation & UI", 
         "• Responsive HTML5 & Modern Vanilla CSS\n• Custom Glassmorphic Dark UI & CSS Grid\n• Interactive JavaScript Timetable & Tabs\n• Chart.js Interactive Analytics Charts",
         COLOR_PRIMARY),
        ("Layer 2: Application & Controllers", 
         "• Python Flask Framework (RESTful Endpoints)\n• Session-Based Role Authentication (RBAC)\n• Secure File Ingestion (PyPDF2, python-docx)\n• Business Logic & Application Status Routing",
         COLOR_ACCENT),
        ("Layer 3: AI & Machine Learning Core", 
         "• Random Forest Classifier (placement_model.pkl)\n• Real-time 8-Dimensional Feature Inference\n• Scikit-Learn TF-IDF Vectorizer & Cosine Metric\n• Heuristic Rule-Based Recommendation Engine",
         COLOR_SUCCESS),
        ("Layer 4: Data & Persistence Layer", 
         "• SQLite Relational Database (smart_campus.db)\n• 15,000 Records Benchmark Dataset (CSV)\n• Entity Tables: students, companies, apps, admins\n• Secure Server-side Resumes Directory (/uploads)",
         COLOR_AMBER)
    ]
    l_w = Inches(2.76)
    l_gap = Inches(0.23)
    for i, (lname, ldesc, lcol) in enumerate(layers):
        lx = Inches(0.8) + i * (l_w + l_gap)
        add_card(slide5, lx, Inches(1.6), l_w, Inches(5.2), CARD_DARK, lcol)

        box = slide5.shapes.add_textbox(lx + Inches(0.2), Inches(1.85), l_w - Inches(0.4), Inches(4.7))
        btf = box.text_frame
        btf.word_wrap = True
        bp0 = btf.paragraphs[0]
        bp0.text = lname
        bp0.font.name = "Segoe UI"
        bp0.font.size = Pt(14)
        bp0.font.bold = True
        bp0.font.color.rgb = lcol
        bp0.space_after = Pt(14)

        for line in ldesc.split("\n"):
            p = btf.add_paragraph()
            p.text = line
            p.font.name = "Segoe UI"
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_MUTED
            p.space_before = Pt(6)

    # =========================================================================
    # SLIDE 6: TECHNOLOGY STACK & TOOLING
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "Comprehensive Technology Stack", "05. IMPLEMENTATION")

    stacks = [
        ("Backend & Web Framework", [
            ("Python 3.10+", "Core backend execution language"),
            ("Flask Framework", "Lightweight, robust WSGI web server"),
            ("Werkzeug", "Secure file handling and utility helpers"),
            ("Jinja2 Engine", "Dynamic server-side HTML template rendering")
        ], COLOR_PRIMARY),
        ("Machine Learning & Data", [
            ("Scikit-Learn", "RandomForestClassifier, metrics, TF-IDF"),
            ("Pandas & NumPy", "Dataframe manipulation & matrix math"),
            ("Joblib", "Fast binary model serialization (.pkl)"),
            ("Matplotlib", "Analytical charts & visual reporting")
        ], COLOR_SUCCESS),
        ("NLP & Text Processing", [
            ("PyPDF2", "Fast programmatic PDF text extraction"),
            ("python-docx", "Microsoft Word (.docx) document parsing"),
            ("Cosine Similarity", "Vectorized textual semantic alignment"),
            ("RegEx Catalogs", "Curated technical & soft skill matchers")
        ], COLOR_PURPLE),
        ("Frontend & Persistence", [
            ("Vanilla CSS3", "Ultra-responsive flexbox & CSS grid"),
            ("Modern JavaScript", "Asynchronous fetch API & live DOM"),
            ("SQLite3 Database", "Embedded zero-config relational store"),
            ("Font Awesome", "Intuitive visual icons & UI cues")
        ], COLOR_AMBER)
    ]
    sw = Inches(2.76)
    sg = Inches(0.23)
    for i, (cat_name, items, col) in enumerate(stacks):
        sx = Inches(0.8) + i * (sw + sg)
        add_card(slide6, sx, Inches(1.6), sw, Inches(5.2), CARD_DARK, CARD_BORDER)

        # Header box
        box = slide6.shapes.add_textbox(sx + Inches(0.2), Inches(1.8), sw - Inches(0.4), Inches(4.8))
        btf = box.text_frame
        btf.word_wrap = True
        
        hp = btf.paragraphs[0]
        hp.text = cat_name
        hp.font.name = "Segoe UI"
        hp.font.size = Pt(13)
        hp.font.bold = True
        hp.font.color.rgb = col
        hp.space_after = Pt(14)

        for name, sub in items:
            p_name = btf.add_paragraph()
            p_name.text = f"• {name}"
            p_name.font.name = "Segoe UI"
            p_name.font.size = Pt(11)
            p_name.font.bold = True
            p_name.font.color.rgb = TEXT_WHITE
            p_name.space_before = Pt(8)

            p_sub = btf.add_paragraph()
            p_sub.text = f"   {sub}"
            p_sub.font.name = "Segoe UI"
            p_sub.font.size = Pt(9.5)
            p_sub.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 7: STUDENT ERP & ACADEMIC MODULES
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "Student ERP & Daily Academic Operations", "06. CORE MODULES")

    modules = [
        ("📅 Dynamic Timetable Engine", 
         "Displays weekly schedules from Monday through Friday with exact class time slots, room numbers, lecture types, and assigned professors. Features live class status indicators and RESTful API (/api/timetable/<day>)."),
        ("📊 Subject-Wise Attendance Tracker", 
         "Tracks real-time attendance across 11 accredited engineering subjects. Computes classes attended versus total conducted sessions. Provides visual progress bars and warning indicators when falling below the mandatory 75% cutoff."),
        ("👨‍🏫 Faculty Directory & Course Hub", 
         "Presents a complete registry of departmental professors, specialized course codes (e.g., BCS501, BCS504, BCS506), syllabus links, and behavioral/open elective details."),
        ("👤 Comprehensive Student Profile Ledger", 
         "Maintains complete student master data including Enrollment ID, Semester, Program, CGPA, Aptitude/Coding scores, Fee clearance ledger, and state/college tier categorization.")
    ]
    for i, (m_title, m_desc) in enumerate(modules):
        my = Inches(1.6) + i * Inches(1.3)
        add_card(slide7, Inches(0.8), my, Inches(11.73), Inches(1.15), CARD_DARK, CARD_BORDER)

        box = slide7.shapes.add_textbox(Inches(1.1), my + Inches(0.12), Inches(11.1), Inches(0.95))
        btf = box.text_frame
        btf.word_wrap = True

        p0 = btf.paragraphs[0]
        p0.text = m_title
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = TEXT_ACCENT

        p1 = btf.add_paragraph()
        p1.text = m_desc
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(10.5)
        p1.font.color.rgb = TEXT_MUTED
        p1.space_before = Pt(4)

    # =========================================================================
    # SLIDE 8: MACHINE LEARNING PLACEMENT PREDICTION ENGINE
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "Machine Learning Placement Prediction Engine", "07. AI/ML SYSTEM")

    # Left Card: Algorithm & Workflow
    add_card(slide8, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), CARD_DARK, COLOR_PRIMARY)
    lbox = slide8.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(5.0), Inches(4.7))
    ltf = lbox.text_frame
    ltf.word_wrap = True

    lp0 = ltf.paragraphs[0]
    lp0.text = "Random Forest Classification Model"
    lp0.font.name = "Segoe UI"
    lp0.font.size = Pt(16)
    lp0.font.bold = True
    lp0.font.color.rgb = TEXT_WHITE

    bullets = [
        "Ensemble Architecture: 100 decision trees (n_estimators=100) with max_depth=12 to prevent overfitting.",
        "Input Vector: 8 quantitative parameters encompassing academic, technical, problem-solving, and interpersonal facets.",
        "Class Probabilities: Generates calibrated percentage probability of placement via predict_proba().",
        "Explainable AI Advice: Dynamic rule engine evaluates student input against benchmarks to produce targeted improvement recommendations.",
        "Model Persistence: Serialized into placement_model.pkl for microsecond server-side inference."
    ]
    for b in bullets:
        p = ltf.add_paragraph()
        p.text = f"• {b}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(8)

    # Right Card: The 8 Predictive Dimensions
    add_card(slide8, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), CARD_DARK, COLOR_SUCCESS)
    rbox = slide8.shapes.add_textbox(Inches(7.1), Inches(1.85), Inches(5.1), Inches(4.7))
    rtf = rbox.text_frame
    rtf.word_wrap = True

    rp0 = rtf.paragraphs[0]
    rp0.text = "The 8 Predictive Feature Dimensions"
    rp0.font.name = "Segoe UI"
    rp0.font.size = Pt(16)
    rp0.font.bold = True
    rp0.font.color.rgb = TEXT_WHITE

    feats = [
        ("1. DSA Problems Solved", "Total coding questions completed on LeetCode/GFG"),
        ("2. CGPA (Cumulative GPA)", "Overall academic score index (scale 0 - 10)"),
        ("3. Communication Score", "Verbal & written score for HR rounds (0 - 100)"),
        ("4. Attendance Percentage", "Regularity & institutional discipline (0 - 100%)"),
        ("5. Aptitude Score", "Quantitative, logical, and verbal screening test (0 - 100)"),
        ("6. Internships Count", "Practical industry internship experience (0 - 4+)"),
        ("7. Projects Count", "Verified full-stack/AI software capstones (0 - 5+)"),
        ("8. Certifications Count", "Industry-recognized technical credentials (0 - 6+)")
    ]
    for f_name, f_desc in feats:
        p = rtf.add_paragraph()
        p.text = f"{f_name}: "
        p.font.name = "Segoe UI"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = TEXT_ACCENT
        p.space_before = Pt(4)

        p2 = rtf.add_paragraph()
        p2.text = f"   {f_desc}"
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 9: DATASET & TRAINING PERFORMANCE
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9)
    add_header(slide9, "Dataset Analysis & Training Evaluation", "08. EMPIRICAL RESULTS")

    # Left: Stats & Confusion Matrix
    add_card(slide9, Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.2), CARD_DARK, CARD_BORDER)
    cm_path = "report_assets/confusion_matrix.png"
    if os.path.exists(cm_path):
        slide9.shapes.add_picture(cm_path, Inches(1.2), Inches(1.9), width=Inches(5.0))

    # Right: Metric cards & Breakdown
    add_card(slide9, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.2), CARD_DARK, CARD_BORDER)
    r_box = slide9.shapes.add_textbox(Inches(7.2), Inches(1.85), Inches(5.0), Inches(4.7))
    rtf = r_box.text_frame
    rtf.word_wrap = True

    rp0 = rtf.paragraphs[0]
    rp0.text = "Model Evaluation Breakdown"
    rp0.font.name = "Segoe UI"
    rp0.font.size = Pt(16)
    rp0.font.bold = True
    rp0.font.color.rgb = TEXT_WHITE
    rp0.space_after = Pt(10)

    perf_points = [
        ("Benchmark Dataset", "Indian Engineering Placement Dataset 2026 (15,000 records)"),
        ("Data Partitioning", "80% Training (12,000) : 20% Testing (3,000 records)"),
        ("Training Accuracy", "89.43% across all classes"),
        ("Test Accuracy", "76.27% on unseen student test records"),
        ("Placed Class Precision", "79.49% (High confidence on positive predictions)"),
        ("Placed Class Recall", "85.67% (Identifies 1,674 of 1,954 placed candidates)"),
        ("Placed Class F1-Score", "82.46% (Strong harmonic balance)")
    ]
    for label, val in perf_points:
        p1 = rtf.add_paragraph()
        p1.text = f"• {label}: "
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(10.5)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_ACCENT
        p1.space_before = Pt(5)

        p2 = rtf.add_paragraph()
        p2.text = f"   {val}"
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 10: FEATURE IMPORTANCE ANALYSIS
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10)
    add_header(slide10, "Feature Importance & Key Placement Drivers", "09. ML EXPLAINABILITY")

    # Left: Feature Importance Chart
    add_card(slide10, Inches(0.8), Inches(1.6), Inches(7.4), Inches(5.2), CARD_DARK, CARD_BORDER)
    fi_path = "report_assets/feature_importance.png"
    if os.path.exists(fi_path):
        slide10.shapes.add_picture(fi_path, Inches(1.0), Inches(1.8), width=Inches(7.0))

    # Right: Analytical Insights
    add_card(slide10, Inches(8.5), Inches(1.6), Inches(4.0), Inches(5.2), CARD_DARK, COLOR_ACCENT)
    ibox = slide10.shapes.add_textbox(Inches(8.75), Inches(1.85), Inches(3.5), Inches(4.7))
    itf = ibox.text_frame
    itf.word_wrap = True

    ip0 = itf.paragraphs[0]
    ip0.text = "Critical Takeaways"
    ip0.font.name = "Segoe UI"
    ip0.font.size = Pt(15)
    ip0.font.bold = True
    ip0.font.color.rgb = TEXT_WHITE
    ip0.space_after = Pt(8)

    insights = [
        "1. DSA is King (34.1%): Solving Data Structures & Algorithm problems is by far the single biggest driver of campus placement.",
        "2. CGPA is the Filter (16.0%): Academic grades serve as the primary cutoff threshold for company shortlisting.",
        "3. Soft Skills Matter (11.2%): Strong communication distinguishes technically proficient students in HR interviews.",
        "4. Consistency Counts: Attendance (9.7%) and Aptitude (9.2%) reflect student commitment and cognitive problem solving."
    ]
    for ins in insights:
        p = itf.add_paragraph()
        p.text = ins
        p.font.name = "Segoe UI"
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 11: INTELLIGENT ATS RESUME ANALYZER
    # =========================================================================
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11)
    add_header(slide11, "Intelligent ATS Resume Analyzer (NLP Engine)", "10. NLP SYSTEM")

    ats_cards = [
        ("Multi-Format Document Parser", 
         "• Extracts clean text from PDF (PyPDF2) and DOCX (python-docx).\n• Sanitizes formatting, normalizes casing, and handles multi-page layouts.\n• Auto-associates uploaded files with student session.", COLOR_PRIMARY),
        ("Curated Skill Catalog Engine", 
         "• Tech Catalog: 64+ languages, frameworks, AI/ML tools, and databases.\n• Soft Skills: 16+ behavioral traits (Leadership, Teamwork, etc.).\n• Action Verbs: 21+ impactful verbs (Architected, Spearheaded, Optimized).", COLOR_SUCCESS),
        ("TF-IDF & Cosine Alignment", 
         "• Computes n-gram TF-IDF vector embeddings for resume and job description.\n• Calculates mathematical cosine similarity metric (0% to 100%).\n• Provides high-precision relevance scores across multiple company profiles.", COLOR_PURPLE),
        ("Automated Gap & ATS Scoring", 
         "• Composite ATS Index: Weighted score of Skills (50%), Similarity (30%), Formatting (10%), and Action Verbs (10%).\n• Actionable feedback: Lists critical missing keywords needed for recruiter shortlisting.", COLOR_AMBER)
    ]
    aw = Inches(2.76)
    ag = Inches(0.23)
    for i, (title, desc, col) in enumerate(ats_cards):
        ax = Inches(0.8) + i * (aw + ag)
        add_card(slide11, ax, Inches(1.6), aw, Inches(5.2), CARD_DARK, col)

        box = slide11.shapes.add_textbox(ax + Inches(0.2), Inches(1.85), aw - Inches(0.4), Inches(4.7))
        btf = box.text_frame
        btf.word_wrap = True

        p0 = btf.paragraphs[0]
        p0.text = title
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = col
        p0.space_after = Pt(12)

        for line in desc.split("\n"):
            p = btf.add_paragraph()
            p.text = line
            p.font.name = "Segoe UI"
            p.font.size = Pt(10)
            p.font.color.rgb = TEXT_MUTED
            p.space_before = Pt(5)

    # =========================================================================
    # SLIDE 12: CAMPUS RECRUITMENT & APPLICATION TRACKING
    # =========================================================================
    slide12 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide12)
    add_header(slide12, "Campus Recruitment & Application Tracking", "11. ATS PORTAL")

    rec_features = [
        ("🏢 Active Corporate Recruitment Drives", 
         "Students explore active on-campus hiring drives from leading firms (e.g., TCS, Infosys, Microsoft, Capgemini, Accenture, Cognizant) with detailed role requirements, CTC packages (up to 18 LPA), locations, and deadlines."),
        ("🛡️ Automated CGPA Cutoff Validation", 
         "System programmatically verifies the candidate's CGPA against the company's minimum threshold (e.g., 8.5 for Microsoft, 7.5 for Infosys). Ineligible applicants are prevented from applying, eliminating screening friction."),
        ("⚡ 1-Click Application Pipeline", 
         "Eligible students submit applications with a single click. The backend registers the timestamped record in the applications table, preventing duplicate submissions."),
        ("📋 Live Status Lifecycle Tracking", 
         "Applicants track real-time hiring progress across multiple stages: 'Applied' ➔ 'Shortlisted' ➔ 'Interview' ➔ 'Selected / Placed', accompanied by official coordinator notes.")
    ]
    for i, (f_title, f_desc) in enumerate(rec_features):
        fy = Inches(1.6) + i * Inches(1.3)
        add_card(slide12, Inches(0.8), fy, Inches(11.73), Inches(1.15), CARD_DARK, CARD_BORDER)

        box = slide12.shapes.add_textbox(Inches(1.1), fy + Inches(0.12), Inches(11.1), Inches(0.95))
        btf = box.text_frame
        btf.word_wrap = True

        p0 = btf.paragraphs[0]
        p0.text = f_title
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = TEXT_ACCENT

        p1 = btf.add_paragraph()
        p1.text = f_desc
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(10.5)
        p1.font.color.rgb = TEXT_MUTED
        p1.space_before = Pt(4)

    # =========================================================================
    # SLIDE 13: ADMINISTRATOR & TPO DASHBOARD
    # =========================================================================
    slide13 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide13)
    add_header(slide13, "Placement Officer (TPO) Command Center", "12. ADMIN PORTAL")

    # 4 KPI Top Cards
    kpis = [
        ("15,000", "Total Registered Students", COLOR_PRIMARY),
        ("9,780+", "Total Placed Candidates", COLOR_SUCCESS),
        ("8.45 LPA", "Average Placed Package", COLOR_PURPLE),
        ("18.0 LPA", "Highest On-Campus Package", COLOR_AMBER)
    ]
    kw = Inches(2.76)
    kg = Inches(0.23)
    for i, (num, sub, col) in enumerate(kpis):
        kx = Inches(0.8) + i * (kw + kg)
        add_card(slide13, kx, Inches(1.6), kw, Inches(1.35), CARD_DARK, col)
        
        box = slide13.shapes.add_textbox(kx + Inches(0.2), Inches(1.75), kw - Inches(0.4), Inches(1.1))
        btf = box.text_frame
        bp0 = btf.paragraphs[0]
        bp0.text = num
        bp0.font.name = "Segoe UI"
        bp0.font.size = Pt(24)
        bp0.font.bold = True
        bp0.font.color.rgb = col

        bp1 = btf.add_paragraph()
        bp1.text = sub
        bp1.font.name = "Segoe UI"
        bp1.font.size = Pt(10)
        bp1.font.color.rgb = TEXT_MUTED
        bp1.space_before = Pt(3)

    # Bottom 2 Cards
    # Left: Administrative Features
    add_card(slide13, Inches(0.8), Inches(3.2), Inches(5.8), Inches(3.6), CARD_DARK, CARD_BORDER)
    lbox = slide13.shapes.add_textbox(Inches(1.1), Inches(3.4), Inches(5.2), Inches(3.2))
    ltf = lbox.text_frame
    ltf.word_wrap = True

    lp0 = ltf.paragraphs[0]
    lp0.text = "Administrative Features"
    lp0.font.name = "Segoe UI"
    lp0.font.size = Pt(15)
    lp0.font.bold = True
    lp0.font.color.rgb = TEXT_WHITE

    admin_bullets = [
        "Dynamic Drive Management: Post new recruitment drives with criteria, package, location, and deadlines.",
        "Candidate Status Workflow: Seamlessly promote student applications from Applied to Interview or Placed.",
        "Automated Status Sync: Marking a candidate 'Placed' instantly updates the student master database.",
        "Multi-Dimensional Filtering: Filter thousands of students by Branch, Placement Status, or search by name."
    ]
    for b in admin_bullets:
        p = ltf.add_paragraph()
        p.text = f"• {b}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(5)

    # Right: Departmental & Analytics
    add_card(slide13, Inches(6.9), Inches(3.2), Inches(5.6), Inches(3.6), CARD_DARK, CARD_BORDER)
    rbox = slide13.shapes.add_textbox(Inches(7.2), Inches(3.4), Inches(5.0), Inches(3.2))
    rtf = rbox.text_frame
    rtf.word_wrap = True

    rp0 = rtf.paragraphs[0]
    rp0.text = "Departmental Analytics & Reporting"
    rp0.font.name = "Segoe UI"
    rp0.font.size = Pt(15)
    rp0.font.bold = True
    rp0.font.color.rgb = TEXT_WHITE

    analytics_bullets = [
        "Branch-Wise Performance: Real-time placement rates across CSE, IT, ECE, Mechanical, and Civil engineering.",
        "Package Dispersion: Average and median package tracking per academic department.",
        "Interactive Dataset Explorer: Live search & review of all 15,000 engineering student benchmark records.",
        "Zero Latency: Optimized SQLite index queries guarantee instantaneous dashboard rendering."
    ]
    for b in analytics_bullets:
        p = rtf.add_paragraph()
        p.text = f"• {b}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(5)

    # =========================================================================
    # SLIDE 14: DATABASE SCHEMA & ENTITY ARCHITECTURE
    # =========================================================================
    slide14 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide14)
    add_header(slide14, "Database Architecture & Entity Relationships", "13. DATA PERSISTENCE")

    tables = [
        ("TABLE: students", [
            ("student_id (PK)", "TEXT PRIMARY KEY"),
            ("name, password", "Authentication & profile"),
            ("course, semester", "Degree & academic standing"),
            ("attendance, cgpa", "Academic metrics"),
            ("aptitude_score", "Quantitative capability"),
            ("coding_score", "DSA problems solved proxy"),
            ("projects, internships", "Practical experience counts"),
            ("skills, branch, state", "Demographics & tech stack"),
            ("placement_status", "Placed / Not Placed status"),
            ("package_lpa", "Secured CTC compensation")
        ], COLOR_PRIMARY),
        ("TABLE: companies", [
            ("id (PK)", "INTEGER AUTOINCREMENT"),
            ("name", "Corporate recruiter name"),
            ("role", "Designation (e.g., SWE)"),
            ("package", "CTC string (e.g. 9.5 LPA)"),
            ("min_cgpa", "Eligibility cutoff filter"),
            ("required_skills", "Target tech requirements"),
            ("location, deadline", "Drive logistical details"),
            ("openings", "Number of available seats")
        ], COLOR_SUCCESS),
        ("TABLE: applications", [
            ("id (PK)", "INTEGER AUTOINCREMENT"),
            ("student_id (FK)", "Links to students table"),
            ("company_id (FK)", "Links to companies table"),
            ("status", "Applied / Shortlist / Placed"),
            ("applied_date", "ISO application timestamp"),
            ("notes", "Recruiter feedback & remarks"),
            ("UNIQUE KEY", "Prevents duplicate applications")
        ], COLOR_PURPLE),
        ("TABLE: admins", [
            ("id (PK)", "INTEGER AUTOINCREMENT"),
            ("username", "Admin login handle"),
            ("password", "Authentication credential"),
            ("name", "Full administrative name"),
            ("role", "TPO / Placement Head"),
            ("email", "Official communication mail")
        ], COLOR_AMBER)
    ]
    tw = Inches(2.76)
    tg = Inches(0.23)
    for i, (t_name, cols, col) in enumerate(tables):
        tx = Inches(0.8) + i * (tw + tg)
        add_card(slide14, tx, Inches(1.6), tw, Inches(5.2), CARD_DARK, col)

        box = slide14.shapes.add_textbox(tx + Inches(0.2), Inches(1.8), tw - Inches(0.4), Inches(4.8))
        btf = box.text_frame
        btf.word_wrap = True

        hp = btf.paragraphs[0]
        hp.text = t_name
        hp.font.name = "Segoe UI"
        hp.font.size = Pt(12)
        hp.font.bold = True
        hp.font.color.rgb = col
        hp.space_after = Pt(10)

        for col_name, col_desc in cols:
            p = btf.add_paragraph()
            p.text = f"• {col_name}"
            p.font.name = "Segoe UI"
            p.font.size = Pt(9.5)
            p.font.bold = True
            p.font.color.rgb = TEXT_WHITE
            p.space_before = Pt(3)

            p2 = btf.add_paragraph()
            p2.text = f"   {col_desc}"
            p2.font.name = "Segoe UI"
            p2.font.size = Pt(8.5)
            p2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 15: VALUE PROPOSITION & COMPARATIVE ADVANTAGES
    # =========================================================================
    slide15 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide15)
    add_header(slide15, "Comparative Value & Institutional Impact", "14. VALUE PROPOSITION")

    stakeholders = [
        ("🎓 For Students", [
            "Early Predictive Diagnosis: Discover skill deficiencies in 3rd/4th year rather than during final semester drives.",
            "Personalized Roadmaps: Receive clear, quantitative guidance on CGPA, DSA problem targets, and project requirements.",
            "ATS Resume Optimization: Maximize resume shortlist rates with instantaneous keyword and cosine semantic matching.",
            "Frictionless Campus Drives: Apply to multiple hiring opportunities with verified eligibility in a single click."
        ], COLOR_PRIMARY),
        ("🏛️ For Placement Cell (TPO)", [
            "80% Reduction in Administrative Overhead: Automated CGPA validation eliminates manual verification of spreadsheets.",
            "Real-Time Hiring Pipelines: Instantly monitor student application lifecycles from application to final offer.",
            "Drive Posting & Communication: Broadcast corporate drives with clear criteria and manage interview slots seamlessly.",
            "Institutional Placement Analytics: Generate departmental reports and benchmark student readiness."
        ], COLOR_SUCCESS),
        ("💼 For Corporate Recruiters", [
            "Pre-Vetted Candidate Funnel: Guaranteed eligibility compliance based on real university grades and metrics.",
            "Standardized ATS Profiles: Structured resumes parsed with clear skill taxonomies and action verb ratings.",
            "Accelerated Hiring Turnaround: Conduct campus drives faster with structured student evaluation data."
        ], COLOR_PURPLE)
    ]
    sw = Inches(3.68)
    sg = Inches(0.33)
    for i, (title, points, col) in enumerate(stakeholders):
        sx = Inches(0.8) + i * (sw + sg)
        add_card(slide15, sx, Inches(1.6), sw, Inches(5.2), CARD_DARK, col)

        box = slide15.shapes.add_textbox(sx + Inches(0.25), Inches(1.85), sw - Inches(0.5), Inches(4.7))
        btf = box.text_frame
        btf.word_wrap = True

        p0 = btf.paragraphs[0]
        p0.text = title
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(15)
        p0.font.bold = True
        p0.font.color.rgb = col
        p0.space_after = Pt(12)

        for pt in points:
            p = btf.add_paragraph()
            p.text = f"• {pt}"
            p.font.name = "Segoe UI"
            p.font.size = Pt(10)
            p.font.color.rgb = TEXT_MUTED
            p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 16: FUTURE ENHANCEMENTS, CONCLUSION & Q&A
    # =========================================================================
    slide16 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide16, RGBColor(7, 12, 28))

    # Header
    add_header(slide16, "Future Enhancements & Concluding Summary", "15. CONCLUSION & ROADMAP")

    # Left: Future Enhancements
    add_card(slide16, Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.2), CARD_DARK, COLOR_ACCENT)
    lbox = slide16.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(5.2), Inches(4.7))
    ltf = lbox.text_frame
    ltf.word_wrap = True

    lp0 = ltf.paragraphs[0]
    lp0.text = "Future Project Roadmap"
    lp0.font.name = "Segoe UI"
    lp0.font.size = Pt(16)
    lp0.font.bold = True
    lp0.font.color.rgb = TEXT_WHITE
    lp0.space_after = Pt(8)

    roadmap = [
        "1. LLM-Powered AI Mock Interviewer: Integrate conversational AI with speech-to-text to simulate HR and technical interviews.",
        "2. Automated Offer Letter OCR & Verification: Verify recruiter compensation CTC breakup and authentic employment offers.",
        "3. Native Mobile Application: Develop cross-platform Flutter/React Native mobile app with real-time push notifications.",
        "4. WhatsApp / SMS Notification Gateway: Instant interview alerts and reminder triggers for scheduled drive rounds.",
        "5. Deep Learning Recommendation Engine: Personalized course & LeetCode problem recommendations based on target company."
    ]
    for r in roadmap:
        p = ltf.add_paragraph()
        p.text = f"• {r}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(8)

    # Right: Conclusion & Thank You Card
    add_card(slide16, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.2), CARD_DARK, COLOR_SUCCESS)
    rbox = slide16.shapes.add_textbox(Inches(7.2), Inches(1.85), Inches(5.0), Inches(4.7))
    rtf = rbox.text_frame
    rtf.word_wrap = True

    rp0 = rtf.paragraphs[0]
    rp0.text = "Concluding Takeaway"
    rp0.font.name = "Segoe UI"
    rp0.font.size = Pt(16)
    rp0.font.bold = True
    rp0.font.color.rgb = TEXT_WHITE
    rp0.space_after = Pt(10)

    p_conc = rtf.add_paragraph()
    p_conc.text = "Smart Campus bridges the critical divide between collegiate academic management and student career readiness. By combining high-accuracy Machine Learning (Random Forest 76.3%), NLP Resume Intelligence (TF-IDF & Cosine Similarity), and an end-to-end recruitment tracking portal, the platform transforms static universities into intelligent career hubs."
    p_conc.font.name = "Segoe UI"
    p_conc.font.size = Pt(11)
    p_conc.font.color.rgb = TEXT_MUTED
    p_conc.space_after = Pt(18)

    p_ty = rtf.add_paragraph()
    p_ty.text = "Thank You!"
    p_ty.font.name = "Segoe UI"
    p_ty.font.size = Pt(28)
    p_ty.font.bold = True
    p_ty.font.color.rgb = COLOR_SUCCESS
    p_ty.alignment = PP_ALIGN.CENTER

    p_qa = rtf.add_paragraph()
    p_qa.text = "Questions & Discussion"
    p_qa.font.name = "Segoe UI"
    p_qa.font.size = Pt(14)
    p_qa.font.bold = True
    p_qa.font.color.rgb = TEXT_ACCENT
    p_qa.alignment = PP_ALIGN.CENTER
    p_qa.space_before = Pt(4)

    output_pptx = "Smart_Campus_Presentation.pptx"
    prs.save(output_pptx)
    print(f"[OK] Successfully created '{output_pptx}' with 16 professional widescreen slides.")

if __name__ == "__main__":
    create_presentation()
